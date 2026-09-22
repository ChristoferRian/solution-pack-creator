#!/usr/bin/env python3
"""Solution Pack Document generator — sections 1-4 only.

House style: ber-basis template `templates/house-style-template.docx`
(styles/numbering/theme/style-set diambil dari 'Solution Pack Pussiberad 2027 V1.1.docx').
Isi konten (customer, URS, solution) tetap disusun dari MoM/input user.

Usage: generate.py <input.json> <output.docx>
Exit codes: 0 ok, 2 invalid input, 1 runtime error.
"""
import copy
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml

TEMPLATE = Path(__file__).resolve().parent.parent / 'templates' / 'house-style-template.docx'

# ---- palette (dari house-style file) ----
BODY = RGBColor(0x22, 0x22, 0x22)      # warna teks normal
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FILL_LABEL = 'E9ECEF'                  # label DocControl / 4.1
FILL_HEADER = '4F81BD'                 # header tabel URS & 4.1
FILL_ZEBRA = 'DBE5F1'                  # baris data ganjil
FILL_WHITE = 'FFFFFF'
BORDER = '9FBAD0'                      # border tabel 0.5pt

# ---- style names (house style) ----
ST_TITLE = 'Title'
ST_H1 = 'Heading 1'
ST_H2 = 'Heading 2'
ST_BODY = 'Normal - H2'
ST_BULLET = 'Bullet List - H2'
ST_TBL_ITEM = 'Table - Item'
ST_TBL_DESC = 'Table - Description'

DC_LABELS = [
    'Document Name', 'Document Code', 'Project Code', 'Reference Lead Code',
    'Lead Name', 'Customer Name', 'End User', 'Year & Semester Delivery',
    'Version', 'Document Owner', 'Document Status', 'Branding Rule',
]

TBLPR_SEQ = ['w:tblStyle', 'w:tblpPr', 'w:tblOverlap', 'w:bidiVisual',
             'w:tblStyleRowBandSize', 'w:tblStyleColBandSize', 'w:tblW', 'w:jc',
             'w:tblCellSpacing', 'w:tblInd', 'w:tblBorders', 'w:shd',
             'w:tblLayout', 'w:tblCellMar', 'w:tblLook', 'w:tblCaption',
             'w:tblDescription', 'w:tblPrChange']
TCPR_SEQ = ['w:cnfStyle', 'w:tcW', 'w:gridSpan', 'w:hMerge', 'w:vMerge',
            'w:tcBorders', 'w:shd', 'w:noWrap', 'w:tcMar', 'w:textDirection',
            'w:tcFitText', 'w:vAlign', 'w:hideMark', 'w:tcPrChange']
TRPR_SEQ = ['w:cnfStyle', 'w:divId', 'w:gridBefore', 'w:gridAfter', 'w:wBefore',
            'w:wAfter', 'w:cantSplit', 'w:trHeight', 'w:tblHeader',
            'w:tblCellSpacing', 'w:jc', 'w:hidden', 'w:ins', 'w:del',
            'w:trPrChange']


def _normalize(parent, seq):
    """Re-append children of parent in canonical OOXML schema order."""
    order = {t: i for i, t in enumerate(seq)}
    children = list(parent)
    children.sort(key=lambda el: order.get('w:' + el.tag.split('}')[1], len(seq)))
    for el in children:
        parent.remove(el)
    for el in children:
        parent.append(el)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear')
    s.set(qn('w:color'), 'auto')
    s.set(qn('w:fill'), fill)
    tcPr.append(s)
    _normalize(tcPr, TCPR_SEQ)


def _tbl_borders(table):
    bd = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), BORDER)
        bd.append(e)
    table._tbl.tblPr.append(bd)


def _tbl_cellmar(table, top=70, left=100, bottom=70, right=100):
    cm = OxmlElement('w:tblCellMar')
    for tag, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement('w:' + tag)
        e.set(qn('w:w'), str(v))
        e.set(qn('w:type'), 'dxa')
        cm.append(e)
    table._tbl.tblPr.append(cm)


def _fixed_widths(table, widths):
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn('w:tblW'))
    if old is not None:
        tblPr.remove(old)
    tw = OxmlElement('w:tblW')
    tw.set(qn('w:w'), str(sum(widths)))
    tw.set(qn('w:type'), 'dxa')
    tblPr.append(tw)
    _normalize(tblPr, TBLPR_SEQ)
    table.autofit = False
    grid = table._tbl.find(qn('w:tblGrid'))
    for gc in list(grid):
        grid.remove(gc)
    for w in widths:
        gc = OxmlElement('w:gridCol')
        gc.set(qn('w:w'), str(w))
        grid.append(gc)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Emu(widths[i] * 635)  # 635 EMU per dxa


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    for tag in ('w:cantSplit', 'w:tblHeader'):
        trPr.append(OxmlElement(tag))
    _normalize(trPr, TRPR_SEQ)


def _no_style(table, doc):
    """Tabel house-style tidak pakai table style (border digambar manual)."""
    table.style = doc.styles['Normal Table']


def _run(p, text, size=None, bold=False, color=None, center=False):
    r = p.add_run(text)
    if size:
        r.font.size = Pt(size)
    if bold:
        r.font.bold = True
    if color is not None:
        r.font.color.rgb = color
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return r


def _vcenter(cell):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


# ---- halaman: watermark draft + footer (house style) ----
DRAFT_WATERMARK = 'DRAFT'
DRAFT_FOOTER = 'Internal Draft - S03 Approval Use Only'
FOOTER_COLOR = RGBColor(0x5A, 0x5A, 0x5A)

_WM_NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
          'xmlns:v="urn:schemas-microsoft-com:vml" '
          'xmlns:o="urn:schemas-microsoft-com:office:office" '
          'xmlns:w10="urn:schemas-microsoft-com:office:word"')


def _watermark_pict(text):
    """VML watermark teks (diagonal, abu-abu) — sama seperti dokumen acuan V1.2."""
    return (
        f'<w:pict {_WM_NS}>'
        '<v:shape id="PowerPlusWaterMarkObject1" o:spid="_x0000_s2049" o:spt="136" '
        'type="#_x0000_t136" '
        'style="position:absolute;left:0pt;height:168.8pt;width:418.45pt;'
        'mso-position-horizontal:center;mso-position-horizontal-relative:margin;'
        'mso-position-vertical:center;mso-position-vertical-relative:margin;'
        'rotation:-2949120f;z-index:-251657216;'
        'mso-width-relative:page;mso-height-relative:page;" '
        'fillcolor="#C0C0C0" filled="t" stroked="f" coordsize="21600,21600" adj="10800">'
        '<v:path/><v:fill on="t" focussize="0,0"/><v:stroke on="f"/>'
        '<v:imagedata o:title=""/>'
        '<o:lock v:ext="edit" aspectratio="t"/>'
        f'<v:textpath on="t" fitshape="t" fitpath="t" trim="t" xscale="f" string="{text}" '
        'style="font-family:Arial Unicode MS;font-size:36pt;'
        'v-same-letter-heights:f;v-text-align:center;"/>'
        '</v:shape></w:pict>'
    )


def build_page_furniture(doc, meta):
    """Watermark + footer di setiap halaman.

    Default house style: watermark teks "DRAFT" (diagonal, abu-abu #C0C0C0) di header
    dan footer "Internal Draft - S03 Approval Use Only" (8pt, #5A5A5A, rata tengah).
    Bisa di-override lewat meta.watermark_text / meta.footer_text (string kosong = tanpa).
    """
    wm_text = meta.get('watermark_text', DRAFT_WATERMARK)
    footer_text = meta.get('footer_text', DRAFT_FOOTER)

    hdr = doc.sections[0].header
    hdr.is_linked_to_previous = False
    hp = hdr.paragraphs[0]
    hp.style = doc.styles['Header']
    if wm_text:
        run = hp.add_run()
        rPr = run._r.get_or_add_rPr()
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '18')
        rPr.append(sz)
        run._r.append(parse_xml(_watermark_pict(wm_text)))

    ftr = doc.sections[0].footer
    ftr.is_linked_to_previous = False
    # tiga paragraf: kosong - teks - kosong (meniru dokumen acuan)
    for i, para in enumerate(ftr.paragraphs):
        if i > 2:
            para._p.getparent().remove(para._p)
    while len(ftr.paragraphs) < 3:
        ftr.add_paragraph()
    for idx, para in enumerate(ftr.paragraphs[:3]):
        para.style = doc.styles['Footer']
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if idx == 1 and footer_text:
            _run(para, footer_text, size=8, color=FOOTER_COLOR)


def _new_page(doc):
    """Paragraf kosong berisi manual page break.

    House style: setiap Heading 1 (kecuali section pertama) selalu mulai di awal
    halaman baru — mengikuti pola di file acuan 'Solution Pack Pussiberad 2027 V1.1'
    (paragraf kosong + <w:br w:type="page"/> sebelum H1, bukan pageBreakBefore di style).
    """
    p = doc.add_paragraph()
    run = p.add_run()
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    run._r.append(br)
    return p


# ---- style helpers ----

def _find_numid(doc, fmt='decimal', preferred=None):
    """Cari numId yang level-0-nya berformat tertentu (mis. decimal untuk kolom No URS).
    `preferred` dipakai kalau numId itu memang ada & formatnya cocok (biar sama dengan file acuan)."""
    numbering = doc.part.numbering_part.element
    abstract = {}
    for a in numbering.findall(qn('w:abstractNum')):
        lvl0 = a.find(qn('w:lvl'))
        if lvl0 is None:
            continue
        f = lvl0.find(qn('w:numFmt'))
        if f is not None:
            abstract[a.get(qn('w:abstractNumId'))] = f.get(qn('w:val'))
    mapping = {}
    for n in numbering.findall(qn('w:num')):
        aid = n.find(qn('w:abstractNumId'))
        if aid is not None:
            mapping[int(n.get(qn('w:numId')))] = abstract.get(aid.get(qn('w:val')))
    if preferred is not None and mapping.get(preferred) == fmt:
        return preferred
    for nid in sorted(mapping):
        if mapping[nid] == fmt:
            return nid
    return None


def _ensure_style(doc, name, make):
    if name in [s.name for s in doc.styles]:
        return doc.styles[name]
    st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles['Normal']
    st.hidden = False
    st.quick_style = False
    make(st)
    return st


def _set_ind(style_or_par, left=None, left_chars=None, hanging=None, hanging_chars=None,
             first_line=None):
    el = style_or_par.element if hasattr(style_or_par, 'element') else style_or_par._p
    pPr = el.get_or_add_pPr() if hasattr(el, 'get_or_add_pPr') else el.find(qn('w:pPr'))
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    if left is not None:
        ind.set(qn('w:left'), str(left))
    if left_chars is not None:
        ind.set(qn('w:leftChars'), str(left_chars))
    if hanging is not None:
        ind.set(qn('w:hanging'), str(hanging))
    if hanging_chars is not None:
        ind.set(qn('w:hangingChars'), str(hanging_chars))
    if first_line is not None:
        ind.set(qn('w:firstLine'), str(first_line))


def ensure_house_styles(doc):
    """Buat style yang dipakai house-style tapi belum ada di template."""
    def mk_bullet(st):
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = 1.05
        _set_ind(st, left=533, left_chars=400, hanging=133, hanging_chars=100)
        src = doc.styles['List Bullet'].element
        numpr = src.find(qn('w:pPr')).find(qn('w:numPr'))
        if numpr is not None:
            st.element.get_or_add_pPr().append(copy.deepcopy(numpr))

    def mk_desc(st):
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.line_spacing = 1.05
        st.paragraph_format.left_indent = Pt(0)
        st.paragraph_format.first_line_indent = Pt(0)
        _set_ind(st, left=0, first_line=0)

    _ensure_style(doc, ST_BULLET, mk_bullet)
    _ensure_style(doc, ST_TBL_DESC, mk_desc)


# ---- validation ----

def _as_list(v):
    if isinstance(v, str):
        return [v]
    return v


def validate(data):
    errs = []
    meta = data.get('meta') or {}
    for k in ('title', 'document_code', 'project_code', 'reference_lead_code',
              'lead_name', 'customer_name', 'end_user', 'year_semester_delivery',
              'version', 'document_owner', 'document_status', 'branding_rule'):
        if k not in meta:
            errs.append(f'meta.{k} missing')
    bg = data.get('background') or {}
    for k in ('customer_background', 'problem_statement', 'objectives'):
        if k not in bg:
            errs.append(f'background.{k} missing')
    if 'problem_statement' in bg and not _as_list(bg['problem_statement']):
        errs.append('background.problem_statement empty')
    if 'objectives' in bg and not _as_list(bg['objectives']):
        errs.append('background.objectives empty')
    urs = data.get('urs')
    if not isinstance(urs, list) or not urs:
        errs.append('urs must be a non-empty list')
    else:
        for i, item in enumerate(urs):
            if not isinstance(item, dict) or not item.get('name') or not item.get('description'):
                errs.append(f'urs[{i}] needs name & description')
    sol = data.get('solution') or {}
    for k in ('name', 'category', 'deployment_context', 'overview'):
        if k not in sol:
            errs.append(f'solution.{k} missing')
    if 'overview' in sol and not _as_list(sol['overview']):
        errs.append('solution.overview empty')

    dc, pc = meta.get('document_code', ''), meta.get('project_code', '')
    if dc and not dc.startswith('SPD-'):
        errs.append(f'document_code must start with SPD- (got: {dc})')
    if pc and not pc.startswith('PRJ-'):
        errs.append(f'project_code must start with PRJ- (got: {pc})')
    if dc.startswith('SPD-') and pc.startswith('PRJ-') and dc[4:] != pc[4:]:
        errs.append(f'project_code must mirror document_code (SPD-/PRJ- suffix mismatch: {dc} vs {pc})')
    if meta.get('version') and not re.match(r'^V\d+\.\d+', meta['version']):
        errs.append(f'version must match V<x.y> + optional changelog (got: {meta["version"]})')
    return errs


# ---- builders ----

def build_title(doc, data):
    """Judul dokumen: HANYA baris 'SOLUTION PACK DOCUMENT'.

    Tidak ada baris judul pengadaan (keputusan Chris, house style v1.1). `meta.title`
    tetap dipakai untuk document properties saja.
    """
    doc.add_paragraph('SOLUTION PACK DOCUMENT', style=ST_TITLE)


def build_doc_control(doc, meta):
    doc.add_paragraph('1. Document Control', style=ST_H1)
    values = [
        ('Document Name', 'Solution Pack Document'),
        ('Document Code', meta['document_code']),
        ('Project Code', meta['project_code']),
        ('Reference Lead Code', meta['reference_lead_code']),
        ('Lead Name', meta['lead_name']),
        ('Customer Name', meta['customer_name']),
        ('End User', meta['end_user']),
        ('Year & Semester Delivery', meta['year_semester_delivery']),
        ('Version', meta['version']),
        ('Document Owner', meta['document_owner']),
        ('Document Status', meta['document_status']),
        ('Branding Rule', meta['branding_rule']),
    ]
    t = doc.add_table(rows=len(values), cols=2)
    _no_style(t, doc)
    for i, (label, val) in enumerate(values):
        c0, c1 = t.rows[i].cells
        shade(c0, FILL_LABEL)
        _run(c0.paragraphs[0], label, size=10, bold=True, color=BODY)
        _run(c1.paragraphs[0], val, size=10, color=BODY)
    _tbl_borders(t)
    _tbl_cellmar(t)
    _fixed_widths(t, [3090, 7278])


def build_background(doc, bg):
    _new_page(doc)
    doc.add_paragraph('2. Project Background', style=ST_H1)
    doc.add_paragraph('2.1 Customer Background', style=ST_H2)
    for para in _as_list(bg['customer_background']):
        doc.add_paragraph(para, style=ST_BODY)
    doc.add_paragraph('2.2 Problem Statement', style=ST_H2)
    for para in _as_list(bg['problem_statement']):
        doc.add_paragraph(para, style=ST_BODY)
    doc.add_paragraph('2.3 Objective', style=ST_H2)
    for o in _as_list(bg['objectives']):
        doc.add_paragraph(o, style=ST_BULLET)


def build_urs(doc, urs):
    _new_page(doc)
    doc.add_paragraph('3. User Requirement Summary', style=ST_H1)
    numid = _find_numid(doc, 'decimal', preferred=7)
    t = doc.add_table(rows=1, cols=3)
    _no_style(t, doc)
    hdr = t.rows[0]
    _repeat_header(hdr)
    for i, label in enumerate(('No', 'User Requirement', 'Deskripsi')):
        c = hdr.cells[i]
        shade(c, FILL_HEADER)
        _vcenter(c)
        _run(c.paragraphs[0], label, size=10, bold=True, color=WHITE, center=True)
    for i, item in enumerate(urs):
        row = t.add_row()
        fill = FILL_ZEBRA if i % 2 == 0 else FILL_WHITE
        for c in row.cells:
            shade(c, fill)
            _vcenter(c)
        # kolom No: auto-numbering (numPr) sesuai house-style
        p_no = row.cells[0].paragraphs[0]
        if numid is not None:
            pPr = p_no._p.get_or_add_pPr()
            numPr = OxmlElement('w:numPr')
            ilvl = OxmlElement('w:ilvl'); ilvl.set(qn('w:val'), '0')
            nid = OxmlElement('w:numId'); nid.set(qn('w:val'), str(numid))
            numPr.append(ilvl); numPr.append(nid)
            pPr.append(numPr)
            _normalize(pPr, ['w:pStyle', 'w:keepNext', 'w:keepLines', 'w:pageBreakBefore',
                             'w:framePr', 'w:widowControl', 'w:numPr', 'w:suppressLineNumbers',
                             'w:pBdr', 'w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku',
                             'w:wordWrap', 'w:overflowPunct', 'w:topLinePunct', 'w:autoSpaceDE',
                             'w:autoSpaceDN', 'w:bidi', 'w:adjustRightInd', 'w:snapToGrid',
                             'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents',
                             'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment',
                             'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle',
                             'w:rPr', 'w:sectPr', 'w:pPrChange'])
        else:
            _run(p_no, str(i + 1), size=10)
        row.cells[1].paragraphs[0].style = doc.styles[ST_TBL_ITEM]
        _run(row.cells[1].paragraphs[0], item['name'])
        row.cells[2].paragraphs[0].style = doc.styles[ST_TBL_DESC]
        _run(row.cells[2].paragraphs[0], item['description'])
    _tbl_borders(t)
    _tbl_cellmar(t)
    _fixed_widths(t, [605, 2370, 7393])


def build_solution(doc, sol):
    _new_page(doc)
    doc.add_paragraph('4. Proposed Solution', style=ST_H1)
    doc.add_paragraph('4.1 Solution Name', style=ST_H2)
    rows = [
        ('Item', 'Deskripsi'),
        ('Solution Name', sol['name']),
        ('Solution Category', sol['category']),
        ('Deployment Context', sol['deployment_context']),
    ]
    t = doc.add_table(rows=4, cols=2)
    _no_style(t, doc)
    for i, label in enumerate(rows[0]):
        c = t.rows[0].cells[i]
        shade(c, FILL_HEADER)
        _vcenter(c)
        p = c.paragraphs[0]
        p.style = doc.styles[ST_TBL_ITEM]
        _run(p, label, color=WHITE, center=True)
    for j, (label, val) in enumerate(rows[1:], start=1):
        c0, c1 = t.rows[j].cells
        shade(c0, FILL_LABEL)
        _vcenter(c0); _vcenter(c1)
        c0.paragraphs[0].style = doc.styles[ST_TBL_ITEM]
        _run(c0.paragraphs[0], label)
        c1.paragraphs[0].style = doc.styles[ST_TBL_DESC]
        _run(c1.paragraphs[0], val)
    _tbl_borders(t)
    _tbl_cellmar(t)
    _fixed_widths(t, [3090, 7278])

    doc.add_paragraph('4.2 Solution Overview', style=ST_H2)
    for para in _as_list(sol['overview']):
        doc.add_paragraph(para, style=ST_BODY)


def main():
    if len(sys.argv) != 3:
        print('usage: generate.py <input.json> <output.docx>', file=sys.stderr)
        return 2
    with open(sys.argv[1], encoding='utf-8') as f:
        data = json.load(f)
    errs = validate(data)
    if errs:
        for e in errs:
            print('INPUT ERROR:', e, file=sys.stderr)
        return 2
    if not TEMPLATE.exists():
        print(f'INPUT ERROR: template tidak ada: {TEMPLATE}', file=sys.stderr)
        return 2
    doc = Document(str(TEMPLATE))
    ensure_house_styles(doc)
    meta = data['meta']
    cp = doc.core_properties
    cp.title = meta.get('title') or 'Solution Pack Document'
    cp.author = meta.get('document_owner') or 'PreSales / PGO'
    cp.last_modified_by = meta.get('document_owner') or 'PreSales / PGO'
    cp.subject = meta.get('customer_name') or ''
    cp.category = 'Solution Pack Document'
    cp.comments = f"{meta.get('document_code','')} | {meta.get('version','')}".strip(' |')
    build_title(doc, data)
    build_page_furniture(doc, meta)
    build_doc_control(doc, meta)
    build_background(doc, data['background'])
    build_urs(doc, data['urs'])
    build_solution(doc, data['solution'])
    doc.save(sys.argv[2])
    print('OK:', sys.argv[2])
    return 0


if __name__ == '__main__':
    sys.exit(main())
