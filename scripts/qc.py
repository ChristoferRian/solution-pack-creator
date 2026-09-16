#!/usr/bin/env python3
"""QC checker for generated Solution Pack .docx (sections 1-4) — house style v1.1.

Usage: qc.py <output.docx>  ->  prints QC PASS or FAIL list; exit 1 on FAIL.
Style acuan: references/style-spec.md (house style dari 'Solution Pack Pussiberad 2027 V1.1').
"""
import re
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn

DC_LABELS = [
    'Document Name', 'Document Code', 'Project Code', 'Reference Lead Code',
    'Lead Name', 'Customer Name', 'End User', 'Year & Semester Delivery',
    'Version', 'Document Owner', 'Document Status', 'Branding Rule',
]
TITLE_TEXT = 'SOLUTION PACK DOCUMENT'
EXPECTED_H1 = ['1. Document Control', '2. Project Background',
               '3. User Requirement Summary', '4. Proposed Solution']
EXPECTED_H2 = ['2.1 Customer Background', '2.2 Problem Statement', '2.3 Objective',
               '4.1 Solution Name', '4.2 Solution Overview']
URS_HEADER = ['No', 'User Requirement', 'Deskripsi']
TBL_HEADER_41 = ['Item', 'Deskripsi']
FORBIDDEN = ['stationery', 'algorhytm', 'one tme', 'spesfikasi', 'anlyzer', 'chipper']

# style house-style yang wajib ada + warna heading
REQUIRED_STYLES = ['Title', 'Heading 1', 'Heading 2', 'Normal', 'Normal - H2',
                   'Bullet List - H2', 'Table - Item', 'Table - Description']
HEADING_COLORS = {'Heading 1': '123555', 'Heading 2': '1F497D'}


def cell_fill(cell):
    tcPr = cell._tc.find(qn('w:tcPr'))
    if tcPr is None:
        return None
    shd = tcPr.find(qn('w:shd'))
    return shd.get(qn('w:fill')) if shd is not None else None


def all_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return '\n'.join(parts)


def styled(doc, level):
    name = f'Heading {level}'
    return [p.text for p in doc.paragraphs if p.style is not None and p.style.name == name]


def has_numpr(p):
    return p._p.find(qn('w:pPr')) is not None and \
        p._p.find(qn('w:pPr')).find(qn('w:numPr')) is not None


def check(doc, errors):
    # --- title + headings ---
    titles = [p.text for p in doc.paragraphs if p.style is not None and p.style.name == 'Title']
    if not titles:
        errors.append('tidak ada paragraf ber-style Title (judul dokumen)')
    for t in titles:
        if t.strip() != TITLE_TEXT:
            errors.append(f'judul tidak sesuai: {t!r} (harus {TITLE_TEXT!r})')

    h1 = styled(doc, 1)
    if h1 != EXPECTED_H1:
        errors.append(f'Heading 1 sequence salah: {h1}')
    h2 = styled(doc, 2)
    if h2 != EXPECTED_H2:
        errors.append(f'Heading 2 sequence salah: {h2}')
    h3 = styled(doc, 3)
    if h3:
        errors.append(f'tidak boleh ada Heading 3 di house style v1.1: {h3}')
    for p in doc.paragraphs:
        if p.style is not None and p.style.name.startswith('Heading') \
                and has_numpr(p):
            errors.append(f'heading pakai auto-numbering (harus manual): {p.text!r}')

    # --- body/bullet paragraph styles ---
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        if p.style is None:
            continue
        if p.style.name == 'Normal' :
            errors.append(f'paragraf body ber-style Normal (harus "Normal - H2"): {p.text[:40]!r}')
        if p.style.name in ('List Bullet', 'List Paragraph'):
            errors.append(f'bullet ber-style {p.style.name!r} (harus "Bullet List - H2"): {p.text[:40]!r}')

    # --- styles wajib ada ---
    names = [s.name for s in doc.styles]
    for s in REQUIRED_STYLES:
        if s not in names:
            errors.append(f'style wajib tidak ada: {s!r}')
    for style_name, want in HEADING_COLORS.items():
        if style_name in names:
            rPr = doc.styles[style_name].element.find(qn('w:rPr'))
            col = rPr.find(qn('w:color')) if rPr is not None else None
            got = col.get(qn('w:val')) if col is not None else None
            if got != want:
                errors.append(f'{style_name} color={got} (harus {want})')

    # --- tabel ---
    tables = doc.tables
    if len(tables) != 3:
        errors.append(f'jumlah tabel harus 3 (DocControl/URS/4.1), dapat {len(tables)}')
    if len(tables) < 3:
        return
    t0, t1, t2 = tables[0], tables[1], tables[2]

    # DocControl
    if len(t0.rows) != 12:
        errors.append(f'DocControl rows: {len(t0.rows)} (harus 12)')
    for i, row in enumerate(t0.rows):
        if i < len(DC_LABELS) and row.cells[0].text != DC_LABELS[i]:
            errors.append(f'DocControl label r{i}: {row.cells[0].text!r} != {DC_LABELS[i]!r}')
        if cell_fill(row.cells[0]) != 'E9ECEF':
            errors.append(f'DocControl label fill r{i}: {cell_fill(row.cells[0])}')
    dc_code = t0.rows[1].cells[1].text
    pr_code = t0.rows[2].cells[1].text
    version = t0.rows[8].cells[1].text
    if not dc_code.startswith('SPD-'):
        errors.append(f'document_code not SPD-: {dc_code!r}')
    if not pr_code.startswith('PRJ-'):
        errors.append(f'project_code not PRJ-: {pr_code!r}')
    if dc_code.startswith('SPD-') and pr_code.startswith('PRJ-') and dc_code[4:] != pr_code[4:]:
        errors.append(f'code mirror mismatch: {dc_code} vs {pr_code}')
    if not re.match(r'^V\d+\.\d+', version):
        errors.append(f'version format wrong: {version!r}')

    # URS
    hdr = [c.text for c in t1.rows[0].cells]
    if hdr != URS_HEADER:
        errors.append(f'URS header wrong: {hdr}')
    for i, c in enumerate(t1.rows[0].cells):
        if cell_fill(c) != '4F81BD':
            errors.append(f'URS header fill c{i}: {cell_fill(c)}')
    for i, row in enumerate(t1.rows[1:]):
        cells = row.cells
        if len(cells) < 3:
            errors.append(f'URS r{i+1}: <3 cells')
            continue
        auto = has_numpr(cells[0].paragraphs[0])
        if not auto and cells[0].text.strip() != str(i + 1):
            errors.append(f'URS No tidak berurutan / tidak auto-numbering di r{i+1}: {cells[0].text!r}')
        expected = 'DBE5F1' if i % 2 == 0 else 'FFFFFF'
        for k in range(3):
            if cell_fill(cells[k]) != expected:
                errors.append(f'URS zebra broken r{i+1} c{k}: {cell_fill(cells[k])} (harus {expected})')
        if cells[1].paragraphs[0].style.name != 'Table - Item':
            errors.append(f'URS r{i+1} nama style={cells[1].paragraphs[0].style.name!r} '
                          f'(harus "Table - Item")')
        if cells[2].paragraphs[0].style.name != 'Table - Description':
            errors.append(f'URS r{i+1} deskripsi style={cells[2].paragraphs[0].style.name!r} '
                          f'(harus "Table - Description")')
        if not cells[1].text.strip() or not cells[2].text.strip():
            errors.append(f'URS r{i+1}: name/description kosong')

    # 4.1
    if len(t2.rows) != 4:
        errors.append(f'4.1 table rows: {len(t2.rows)} (harus 4)')
    else:
        hdr2 = [c.text for c in t2.rows[0].cells]
        if hdr2 != TBL_HEADER_41:
            errors.append(f'4.1 header wrong: {hdr2}')
        for i, c in enumerate(t2.rows[0].cells):
            if cell_fill(c) != '4F81BD':
                errors.append(f'4.1 header fill c{i}: {cell_fill(c)}')
        labels = ['Solution Name', 'Solution Category', 'Deployment Context']
        for j, want in enumerate(labels, start=1):
            if t2.rows[j].cells[0].text != want:
                errors.append(f'4.1 label r{j}: {t2.rows[j].cells[0].text!r} != {want!r}')
            if cell_fill(t2.rows[j].cells[0]) != 'E9ECEF':
                errors.append(f'4.1 label fill r{j}: {cell_fill(t2.rows[j].cells[0])}')

    # baris kosong nyangkut
    for ti, t in enumerate(tables):
        last = t.rows[-1]
        if not any(c.text.strip() for c in last.cells):
            errors.append(f'tabel {ti+1}: ada baris kosong di akhir')

    # typo terlarang
    text = all_text(doc).lower()
    for bad in FORBIDDEN:
        if bad in text:
            errors.append(f'kata terlarang: {bad!r}')


def check_theme(path, errors):
    """Theme harus Calibri (heading) + Cambria (body) sesuai house style."""
    try:
        with zipfile.ZipFile(path) as z:
            theme = z.read('word/theme/theme1.xml').decode('utf-8', 'ignore')
    except Exception as e:
        errors.append(f'tidak bisa baca theme: {e}')
        return
    major = re.search(r'<a:majorFont>.*?<a:latin typeface="([^"]*)"', theme, re.S)
    minor = re.search(r'<a:minorFont>.*?<a:latin typeface="([^"]*)"', theme, re.S)
    if not major or major.group(1) != 'Calibri':
        errors.append(f'theme majorFont bukan Calibri: {major.group(1) if major else None}')
    if not minor or minor.group(1) != 'Cambria':
        errors.append(f'theme minorFont bukan Cambria: {minor.group(1) if minor else None}')


def main():
    if len(sys.argv) != 2:
        print('usage: qc.py <output.docx>', file=sys.stderr)
        return 2
    path = sys.argv[1]
    doc = Document(path)
    errors = []
    check(doc, errors)
    check_theme(path, errors)
    if errors:
        for e in errors:
            print('FAIL:', e)
        print(f'QC FAIL: {len(errors)} issue(s)')
        return 1
    print('QC PASS: no issues')
    return 0


if __name__ == '__main__':
    sys.exit(main())
